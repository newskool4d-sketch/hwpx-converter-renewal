from __future__ import annotations

import base64
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.shared import Mm

import anyway_to_hwpx_com as converter


PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScLbtAAAAABJRU5ErkJggg==')


class DocumentMediaTests(unittest.TestCase):
    def setUp(self):
        converter.pop_conversion_notes()

    def test_markdown_local_image_order_size_and_unicode_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / '그림 파일.png').write_bytes(PNG)
            source = root / '문서.md'
            source.write_text('앞\n\n![그림](<그림 파일.png>)\n\n뒤', encoding='utf-8')
            blocks = converter.detect_and_parse(source)
            self.assertEqual([block['type'] for block in blocks], ['p', 'image', 'p'])
            self.assertEqual(Path(blocks[1]['path']), root / '그림 파일.png')
            self.assertGreater(blocks[1]['width_mm'], 0)
            self.assertEqual(converter.pop_conversion_notes(), [])

    def test_remote_unc_and_missing_markdown_images_do_not_download(self):
        for target in ['https://example.test/image.png', '//server/share/image.png', r'\\server\share\image.png', 'data:image/png;base64,abc', 'missing.png']:
            with self.subTest(target=target), tempfile.TemporaryDirectory() as tmp:
                blocks = converter.parse_markdown(f'앞\n![그림]({target})\n뒤', source_dir=tmp)
                self.assertEqual([block['text'] for block in blocks], ['앞', '뒤'])
                self.assertTrue(any('미삽입' in message for message in converter.pop_conversion_notes()))

    def test_inline_and_code_images_keep_explicit_existing_limits(self):
        with tempfile.TemporaryDirectory() as tmp:
            (Path(tmp) / 'a.png').write_bytes(PNG)
            blocks = converter.parse_markdown('앞 ![그림](a.png) 뒤\n```\n![코드](a.png)\n```', source_dir=tmp)
        self.assertEqual([block['type'] for block in blocks], ['p', 'code'])
        self.assertIn('이미지 1개 미삽입', '\n'.join(converter.pop_conversion_notes()))

    def test_docx_run_text_image_text_and_original_size_survive(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / 'a.png'
            image_path.write_bytes(PNG)
            document = Document()
            paragraph = document.add_paragraph()
            run = paragraph.add_run('앞')
            run.add_picture(str(image_path), width=Mm(30))
            run.add_text('뒤')
            source = root / '문서.docx'
            document.save(source)
            blocks = converter.parse_docx(source, asset_dir=root)
            self.assertEqual([block['type'] for block in blocks], ['p', 'image', 'p'])
            self.assertEqual([blocks[0]['text'], blocks[2]['text']], ['앞', '뒤'])
            self.assertAlmostEqual(blocks[1]['width_mm'], 30)
            self.assertEqual(Path(blocks[1]['path']).read_bytes(), PNG)
            self.assertTrue(any('인라인 이미지' in note for note in converter.pop_conversion_notes()))

    def test_docx_alternate_content_emits_only_one_image(self):
        import copy
        from lxml import etree

        with tempfile.TemporaryDirectory() as tmp:
            image = Path(tmp) / 'a.png'
            image.write_bytes(PNG)
            document = Document()
            run = document.add_paragraph().add_run()
            run.add_picture(str(image))
            drawing = run._r[0]
            run._r.remove(drawing)
            mc = '{http://schemas.openxmlformats.org/markup-compatibility/2006}'
            alternate = etree.SubElement(run._r, mc + 'AlternateContent')
            choice = etree.SubElement(alternate, mc + 'Choice')
            choice.append(drawing)
            fallback = etree.SubElement(alternate, mc + 'Fallback')
            fallback.append(copy.deepcopy(drawing))
            segments = list(converter._docx_paragraph_segments(document.paragraphs[0]))
            self.assertEqual(len(segments), 1)
            self.assertEqual(segments[0]['type'], 'image')

    def test_docx_text_events_match_existing_run_text(self):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        document = Document()
        paragraph = document.add_paragraph()
        run = paragraph.add_run('2026')
        run._r.append(OxmlElement('w:noBreakHyphen'))
        run.add_text('001')
        run._r.append(OxmlElement('w:ptab'))
        run.add_text('뒤')
        page_break = OxmlElement('w:br')
        page_break.set(qn('w:type'), 'page')
        run._r.append(page_break)
        run.add_text('끝')
        self.assertEqual(list(converter._docx_paragraph_segments(paragraph)), [run.text])

    def test_docx_image_without_asset_folder_is_in_memory_not_lost(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / 'a.png'
            image_path.write_bytes(PNG)
            document = Document()
            document.add_picture(str(image_path))
            source = root / 'sample.docx'
            document.save(source)
            blocks = converter.parse_docx(source)
            self.assertEqual(blocks[0]['type'], 'image')
            self.assertEqual(blocks[0]['data'], PNG)

    def test_docx_table_image_has_loss_warning(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image_path = root / 'a.png'
            image_path.write_bytes(PNG)
            document = Document()
            table = document.add_table(rows=1, cols=1)
            paragraph = table.cell(0, 0).paragraphs[0]
            paragraph.add_run('셀 내용').add_picture(str(image_path))
            source = root / 'sample.docx'
            document.save(source)
            blocks = converter.parse_docx(source)
            self.assertEqual(blocks[0]['header'], ['셀 내용'])
            self.assertTrue(any('표 셀 내부 이미지' in note for note in converter.pop_conversion_notes()))

    def test_body_writer_embeds_and_propagates_insert_failure(self):
        calls = []

        class Hwp:
            def InsertPicture(self, *args):
                calls.append(args)
                self.path = Path(args[0])
                self.exists_during_insert = self.path.is_file()
                return False

        hwp = Hwp()
        with self.assertRaises(RuntimeError):
            converter.insert_body_image(hwp, {'type': 'image', 'data': PNG, 'extension': 'png', 'width_mm': 30, 'height_mm': 20})
        self.assertTrue(hwp.exists_during_insert)
        self.assertFalse(hwp.path.exists())
        self.assertEqual(calls[0][1:], (True, 1, False, False, 0, 30.0, 20.0))

    def test_docx_assets_exist_through_save_and_cleanup_after_failure(self):
        class Documents:
            Count = 1
            def Add(self, **kwargs):
                pass
            def Item(self, index):
                return self
            def Close(self, **kwargs):
                pass

        class Hwp:
            XHwpDocuments = Documents()
            def SaveAs(self, *args):
                self.exists_at_save = self.image_path.is_file()
                raise RuntimeError('save failure')

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / 'a.png'
            image.write_bytes(PNG)
            document = Document()
            document.add_picture(str(image))
            source = root / 'sample.docx'
            document.save(source)
            hwp = Hwp()

            def build(hwp, blocks, **kwargs):
                hwp.image_path = Path(blocks[0]['path'])

            with patch.object(converter, 'build_doc', side_effect=build), patch.object(converter.time, 'sleep'):
                with self.assertRaisesRegex(RuntimeError, 'save failure'):
                    converter.convert_file(hwp, source, root / 'sample.hwpx')
            self.assertTrue(hwp.exists_at_save)
            self.assertFalse(hwp.image_path.exists())
            self.assertTrue(source.is_file())


if __name__ == '__main__':
    unittest.main()
